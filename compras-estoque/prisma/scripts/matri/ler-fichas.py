# Lê a planilha de fichas técnicas do Matri e mede o que dá para importar.
#
#   pip install openpyxl
#   python3 -I prisma/scripts/matri/ler-fichas.py <arquivo.xlsx>
#
# NÃO toca no banco. É o passo de medir antes de prometer.
#
# A PLANILHA, como ela é: uma aba por ficha, com duas famílias de nome —
# "MONTAGEM - X" para prato e "SB - X" para sub-receita. Isso casa direto com o
# modelo do app, onde uma receita pode ter outra dentro (IngredienteReceita tem
# produtoId OU subReceitaId). O ragu de ossobuco vira uma receita só, e o nhoque
# e a lasanha puxam dela em vez de repetir os ingredientes nas duas.
#
# O LAYOUT de cada aba:
#   A1 = nome          C1 = "CUSTO DA RECEITA"   E1 = custo    F1 = rendimento
#   linha 2 = cabeçalho: INGREDIENTE | PESO LÍQUIDO | UNIDADE | RENDIMENTO |
#                        PESO BRUTO | PREÇO | CUSTO DO INGREDIENTE
#   linha 3+ = ingredientes
#
# A QUANTIDADE QUE IMPORTA É O **PESO BRUTO**, não o líquido: é ele que
# multiplica o preço para dar o custo (confirmado conferindo custo = bruto x
# preço linha a linha). Peso líquido é o que chega ao prato depois da limpeza;
# peso bruto é o que sai do estoque — e estoque é o que este app mede.
#
# AS ABAS DE APOIO:
#   INSUMOS SISTEMA  — catálogo com CÓDIGO (12 dígitos), o identificador bom
#   CADASTRO NOVOS   — o que ainda não existe, com custo manual
#   LISTA            — as duas consolidadas
#   DASHBOARD        — por prato: categoria, preço de venda e CMV
import re, sys, unicodedata, collections, json
import openpyxl

LIXO = re.compile(r'^(INGREDIENTE|UNIDADE|CUSTO)|^Tempo de forno|^Finalizar|^Massa passada|^Polpettone selado|^\d+g ', re.I)

def chave(s):
    s = unicodedata.normalize("NFD", str(s or "").upper())
    s = "".join(c for c in s if unicodedata.category(c) != "Mn")
    return " ".join(re.sub(r'[^A-Z0-9 ]', ' ', s).split())

def ler(caminho):
    wb = openpyxl.load_workbook(caminho, data_only=True)
    fichas = []
    for nome in wb.sheetnames:
        if not (nome.startswith("MONTAGEM") or nome.startswith("SB")):
            continue
        ws = wb[nome]
        ings = []
        for r in ws.iter_rows(min_row=3, values_only=True):
            ing, liq, un, rend, bruto, preco, custo = (list(r) + [None] * 7)[:7]
            if not ing or not str(ing).strip():
                continue
            # Linha de texto livre no meio da ficha (modo de preparo, tempo de
            # forno) não é ingrediente. Entra no modoPreparo, não na conta.
            if LIXO.match(str(ing)) or str(un or "").upper() == "CUSTO DA RECEITA":
                continue
            ings.append(dict(ing=str(ing).strip(), un=(str(un).strip() if un else None),
                             bruto=bruto, preco=preco, custo=custo))
        fichas.append(dict(aba=nome, titulo=ws["A1"].value, custo=ws["E1"].value,
                           rendimento=ws["F1"].value,
                           tipo="SUB" if nome.startswith("SB") else "PRATO", ings=ings))
    return wb, fichas

if __name__ == "__main__":
    wb, fichas = ler(sys.argv[1])
    pratos = [f for f in fichas if f["tipo"] == "PRATO"]
    subs = [f for f in fichas if f["tipo"] == "SUB"]
    linhas = sum(len(f["ings"]) for f in fichas)
    print(f"{len(pratos)} pratos · {len(subs)} sub-receitas · {linhas} ingredientes")

    # O rendimento é o campo mais perigoso da importação: nulo faz o app tratar
    # a receita como "rende 1", e aí cada prato vendido consome o LOTE INTEIRO.
    # Foi o que aconteceu no Tartare de Atum do Beira Lago, onde um rendimento
    # errado fazia cada prato puxar 0,87 kg de atum.
    sem_rend = [f["aba"] for f in subs if not f["rendimento"]]
    print(f"\nsub-receitas SEM rendimento declarado: {len(sem_rend)}")
    for a in sem_rend:
        print(f"   {a}")

    sist = {chave(r[1]): r[0] for r in wb["INSUMOS SISTEMA"].iter_rows(min_row=4, values_only=True) if r[0] and r[1]}
    novos = {chave(r[0]) for r in wb["CADASTRO NOVOS"].iter_rows(min_row=4, values_only=True) if r[0]}
    titulos = {chave(f["titulo"]) for f in subs} | {chave(f["aba"].replace("SB -", "").replace("SB-", "")) for f in subs}

    onde = collections.Counter()
    orfaos = collections.Counter()
    for f in fichas:
        for i in f["ings"]:
            k = chave(i["ing"])
            if k in titulos: onde["sub-receita desta planilha"] += 1
            elif k in sist: onde["catálogo do sistema (tem código)"] += 1
            elif k in novos: onde["cadastro novo"] += 1
            else:
                onde["sem origem conhecida"] += 1
                orfaos[i["ing"]] += 1
    print()
    for o, n in onde.most_common():
        print(f"  {n:>4}  {o}")
    if orfaos:
        print(f"\nsem origem ({len(orfaos)} nomes):")
        for n, c in orfaos.most_common():
            print(f"   {c}x  {n}")
